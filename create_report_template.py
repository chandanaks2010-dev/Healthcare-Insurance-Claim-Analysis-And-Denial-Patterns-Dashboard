#!/usr/bin/env python
"""
Script to create a professional DOCX project report template
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime

def create_project_report():
    """Create a professional project report document"""
    
    doc = Document()
    
    # Set default style
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    
    # ===== TITLE PAGE =====
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run('Healthcare Insurance Claim Analysis\nAnd Denial Patterns Dashboard')
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph()
    
    # Subtitle
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run('SQL and Tableau Data Visualization Project')
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.italic = True
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Team details section
    team_info = doc.add_paragraph()
    team_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    team_run = team_info.add_run('Project Report Documentation')
    team_run.font.size = Pt(12)
    team_run.font.bold = True
    
    doc.add_paragraph()
    doc.add_paragraph(f'Date: {datetime.now().strftime("%B %d, %Y")}')
    doc.add_paragraph('Course: SQL and Tableau for Data Visualization')
    doc.add_paragraph('Total Marks: 30 marks')
    
    doc.add_page_break()
    
    # ===== TABLE OF CONTENTS =====
    doc.add_heading('Table of Contents', level=1)
    toc_items = [
        '1. Dataset Description',
        '2. Dataset Understanding & Variables',
        '3. SQL Data Preparation & Queries',
        '4. Data Analysis & Aggregations',
        '5. Tableau Visualizations & Dashboard',
        '6. Statistical & KPI Analysis',
        '7. Key Findings & Insights',
        '8. Business Recommendations',
        '9. Conclusion',
        '10. Appendix'
    ]
    for item in toc_items:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_page_break()
    
    # ===== SECTION 1: DATASET DESCRIPTION =====
    doc.add_heading('1. Dataset Description', level=1)
    doc.add_heading('1.1 Overview', level=2)
    doc.add_paragraph(
        'Provide a comprehensive overview of your dataset. Include information about:'
    )
    doc.add_paragraph('Source of the dataset', style='List Bullet')
    doc.add_paragraph('Time period covered', style='List Bullet')
    doc.add_paragraph('Total number of records', style='List Bullet')
    doc.add_paragraph('Key business context', style='List Bullet')
    
    doc.add_paragraph()
    p = doc.add_paragraph('[Add your dataset description here]')
    p_format = p.paragraph_format
    p_format.left_indent = Inches(0.5)
    
    doc.add_heading('1.2 Dataset Characteristics', level=2)
    doc.add_paragraph('File size and format', style='List Bullet')
    doc.add_paragraph('Data collection methodology', style='List Bullet')
    doc.add_paragraph('Data quality observations', style='List Bullet')
    
    doc.add_paragraph()
    p = doc.add_paragraph('[Add dataset characteristics]')
    p_format = p.paragraph_format
    p_format.left_indent = Inches(0.5)
    
    doc.add_page_break()
    
    # ===== SECTION 2: DATASET UNDERSTANDING & VARIABLES =====
    doc.add_heading('2. Dataset Understanding & Variables', level=1)
    doc.add_heading('2.1 Data Dictionary', level=2)
    doc.add_paragraph(
        'Document all columns/variables in your dataset with descriptions and data types:'
    )
    
    # Sample table
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Column Name'
    hdr_cells[1].text = 'Data Type'
    hdr_cells[2].text = 'Description'
    hdr_cells[3].text = 'Business Meaning'
    
    # Add sample row with instructions
    row_cells = table.add_row().cells
    row_cells[0].text = '[e.g., claim_id]'
    row_cells[1].text = '[INT]'
    row_cells[2].text = '[Unique identifier for each claim]'
    row_cells[3].text = '[Primary key]'
    
    doc.add_paragraph('(Replace with your actual columns)')
    
    doc.add_heading('2.2 Dimensions and Measures', level=2)
    doc.add_paragraph('Dimensions (Categorical/Qualitative data):')
    for i in range(3):
        doc.add_paragraph(f'Dimension {i+1}: [Name and description]', style='List Bullet')
    
    doc.add_paragraph('Measures (Quantitative data):')
    for i in range(3):
        doc.add_paragraph(f'Measure {i+1}: [Name and description]', style='List Bullet')
    
    doc.add_page_break()
    
    # ===== SECTION 3: SQL DATA PREPARATION =====
    doc.add_heading('3. SQL Data Preparation & Queries', level=1)
    doc.add_heading('3.1 Data Extraction Strategy', level=2)
    doc.add_paragraph('Describe how you extracted data from the raw dataset')
    doc.add_paragraph('Initial data cleaning and validation steps', style='List Bullet')
    doc.add_paragraph('Data type conversions performed', style='List Bullet')
    doc.add_paragraph('Null value handling', style='List Bullet')
    
    doc.add_heading('3.2 SQL Queries Used', level=2)
    doc.add_paragraph('Provide details of key SQL operations:')
    
    doc.add_heading('SELECT Queries', level=3)
    doc.add_paragraph('[Describe your SELECT queries]')
    
    doc.add_heading('JOINs', level=3)
    doc.add_paragraph('Types of joins used:', style='List Bullet')
    doc.add_paragraph('Tables joined:', style='List Bullet')
    doc.add_paragraph('Join conditions:', style='List Bullet')
    
    doc.add_heading('GROUP BY & HAVING Clauses', level=3)
    doc.add_paragraph('[Describe aggregation logic and filtering]')
    
    doc.add_heading('Subqueries & CTEs', level=3)
    doc.add_paragraph('[Describe complex query logic]')
    
    doc.add_heading('Window Functions', level=3)
    doc.add_paragraph('Functions used (RANK, DENSE_RANK, ROW_NUMBER, LAG, LEAD, etc.):', style='List Bullet')
    doc.add_paragraph('Use cases for each window function:', style='List Bullet')
    
    doc.add_heading('3.3 Data Quality Validation', level=2)
    doc.add_paragraph('[Describe validation checks performed]')
    
    doc.add_page_break()
    
    # ===== SECTION 4: DATA ANALYSIS & AGGREGATIONS =====
    doc.add_heading('4. Data Analysis & Aggregations', level=1)
    doc.add_heading('4.1 Summary Statistics', level=2)
    
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Metric'
    hdr_cells[1].text = 'Value'
    hdr_cells[2].text = 'Unit'
    hdr_cells[3].text = 'Interpretation'
    
    for i in range(4):
        row_cells = table.add_row().cells
        row_cells[0].text = f'[Metric {i+1}]'
        row_cells[1].text = '[Value]'
        row_cells[2].text = '[Unit]'
        row_cells[3].text = '[Business meaning]'
    
    doc.add_heading('4.2 Key Aggregations', level=2)
    doc.add_paragraph('[Describe aggregations by key dimensions]')
    doc.add_paragraph('Total claims by status', style='List Bullet')
    doc.add_paragraph('Average claim amount by category', style='List Bullet')
    doc.add_paragraph('Distribution of denials by reason', style='List Bullet')
    
    doc.add_page_break()
    
    # ===== SECTION 5: TABLEAU VISUALIZATIONS =====
    doc.add_heading('5. Tableau Visualizations & Dashboard', level=1)
    doc.add_heading('5.1 Dashboard Overview', level=2)
    doc.add_paragraph('[Describe your main dashboard and its purpose]')
    
    doc.add_heading('5.2 Worksheets & Visualizations', level=2)
    
    visualizations = [
        ('KPI Dashboard', 'Executive summary with key metrics'),
        ('Trend Analysis', 'Time-series line chart showing patterns'),
        ('Pareto Analysis', 'Bar chart identifying key contributors'),
        ('Geographical Heatmap', 'Regional performance visualization'),
        ('Demographic Analysis', 'Customer segment breakdown'),
        ('RFM Analysis', 'Customer segmentation by Recency, Frequency, Monetary'),
    ]
    
    for viz_name, description in visualizations:
        doc.add_heading(f'  {viz_name}', level=3)
        p = doc.add_paragraph(f'{description}')
        p_format = p.paragraph_format
        p_format.left_indent = Inches(0.5)
        doc.add_paragraph('Chart type: [Type]', style='List Bullet')
        doc.add_paragraph('Dimensions used: [List]', style='List Bullet')
        doc.add_paragraph('Measures aggregated: [List]', style='List Bullet')
        doc.add_paragraph('Interactive filters: [List]', style='List Bullet')
    
    doc.add_heading('5.3 Dashboard Interactivity', level=2)
    doc.add_paragraph('Cascading filters implemented', style='List Bullet')
    doc.add_paragraph('Custom actions configured', style='List Bullet')
    doc.add_paragraph('Drill-down capabilities', style='List Bullet')
    doc.add_paragraph('Tooltip customizations', style='List Bullet')
    
    doc.add_page_break()
    
    # ===== SECTION 6: STATISTICAL & KPI ANALYSIS =====
    doc.add_heading('6. Statistical & KPI Analysis', level=1)
    doc.add_heading('6.1 KPI Definitions', level=2)
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'KPI Name'
    hdr_cells[1].text = 'Formula/Definition'
    hdr_cells[2].text = 'Business Impact'
    
    for i in range(5):
        row_cells = table.add_row().cells
        row_cells[0].text = f'[KPI {i+1}]'
        row_cells[1].text = '[Formula]'
        row_cells[2].text = '[Relevance]'
    
    doc.add_heading('6.2 Trend Analysis', level=2)
    doc.add_paragraph('[Describe trends observed over time]')
    doc.add_paragraph('Seasonal patterns identified', style='List Bullet')
    doc.add_paragraph('Growth or decline trends', style='List Bullet')
    doc.add_paragraph('Forecast predictions', style='List Bullet')
    
    doc.add_heading('6.3 Pareto Analysis (80/20 Rule)', level=2)
    doc.add_paragraph('[Identify key contributors]')
    
    doc.add_heading('6.4 RFM Analysis (if applicable)', level=2)
    doc.add_paragraph('[Segmentation by Recency, Frequency, Monetary value]')
    doc.add_paragraph('Customer segments identified', style='List Bullet')
    doc.add_paragraph('Characteristics of each segment', style='List Bullet')
    doc.add_paragraph('Business actions per segment', style='List Bullet')
    
    doc.add_page_break()
    
    # ===== SECTION 7: KEY FINDINGS & INSIGHTS =====
    doc.add_heading('7. Key Findings & Insights', level=1)
    doc.add_heading('7.1 Major Patterns Discovered', level=2)
    
    for i in range(5):
        doc.add_heading(f'Finding {i+1}', level=3)
        doc.add_paragraph('[Describe the finding]')
        doc.add_paragraph('Data evidence: [Specific numbers/percentages]', style='List Bullet')
        doc.add_paragraph('Visualization reference: [Which chart shows this]', style='List Bullet')
    
    doc.add_heading('7.2 Business Insights', level=2)
    doc.add_paragraph('[What does this mean for the business?]')
    
    doc.add_page_break()
    
    # ===== SECTION 8: BUSINESS RECOMMENDATIONS =====
    doc.add_heading('8. Business Recommendations', level=1)
    
    for i in range(4):
        doc.add_heading(f'Recommendation {i+1}', level=2)
        doc.add_paragraph('[State the recommendation]')
        doc.add_paragraph('Rationale: [Why this action]', style='List Bullet')
        doc.add_paragraph('Expected impact: [Projected outcome]', style='List Bullet')
        doc.add_paragraph('Implementation: [How to execute]', style='List Bullet')
    
    doc.add_page_break()
    
    # ===== SECTION 9: CONCLUSION =====
    doc.add_heading('9. Conclusion', level=1)
    doc.add_paragraph(
        '[Summarize your project findings and the value delivered through this analysis. '
        'Reiterate key insights and how the SQL-Tableau integration provided comprehensive '
        'business intelligence.]'
    )
    
    doc.add_heading('9.1 Project Objectives Achievement', level=2)
    doc.add_paragraph('Objective 1: [Status]', style='List Bullet')
    doc.add_paragraph('Objective 2: [Status]', style='List Bullet')
    doc.add_paragraph('Objective 3: [Status]', style='List Bullet')
    
    doc.add_heading('9.2 Future Enhancements', level=2)
    doc.add_paragraph('[Potential improvements or extensions to this analysis]')
    
    doc.add_page_break()
    
    # ===== SECTION 10: APPENDIX =====
    doc.add_heading('10. Appendix', level=1)
    doc.add_heading('10.1 SQL Scripts', level=2)
    doc.add_paragraph('[Reference to complete SQL scripts in accompanying .sql file]')
    
    doc.add_heading('10.2 Data Samples', level=2)
    doc.add_paragraph('[Sample records from processed data]')
    
    doc.add_heading('10.3 Technical Details', level=2)
    doc.add_paragraph('Tools used: SQL Server / MySQL / PostgreSQL', style='List Bullet')
    doc.add_paragraph('Tableau version: [Version]', style='List Bullet')
    doc.add_paragraph('Data processing framework: [Framework]', style='List Bullet')
    
    doc.add_heading('10.4 References', level=2)
    doc.add_paragraph('Project guidelines document', style='List Bullet')
    doc.add_paragraph('Data source documentation', style='List Bullet')
    doc.add_paragraph('Tableau best practices', style='List Bullet')
    
    # Save document
    doc.save('Project_Report_Documentation.docx')
    print("✓ Project report template created successfully!")
    print("✓ File: Project_Report_Documentation.docx")

if __name__ == '__main__':
    create_project_report()
